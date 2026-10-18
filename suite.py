
import json,sys,re,os,statistics
from playwright.sync_api import sync_playwright
import openpyxl,warnings;warnings.simplefilter('ignore')
HERE=os.path.dirname(os.path.abspath(__file__))
d=json.load(open(os.path.join(HERE,'data.json')));R=d['recs']
MAP=dict(zip('012345678','OIZEASGTB'))
def comp(v): return ''.join(MAP.get(c,c) if c.isdigit() else c for c in v.replace(' ',''))
priced=[r for r in R if r['p'] is not None]
exp=dict(wig=sum(1 for r in R if 'WIG' in comp(r['v'])),boy=sum(1 for r in R if 'BOY' in comp(r['v'])),
 cur=sum(1 for r in R if r['t']=='Current'),cursuf=sum(1 for r in R if r['t'] in('Current','Suffix')),
 p1u250=sum(1 for r in R if r['part']=='Part 1' and r['p'] is not None and r['p']<250),
 cheapest=min(priced,key=lambda r:r['p'])['v'],dearest=max(priced,key=lambda r:r['p'])['v'])
URL='file://'+os.path.abspath(os.path.join(HERE,'..','vrn_dashboard.html'))
dev=sys.argv[1]; heavy=sys.argv[2]=='1'; OUT=sys.argv[3] if len(sys.argv)>3 else HERE
with sync_playwright() as p:
    if dev.startswith('W'):
        w,h=map(int,dev[1:].split('x')); cfg=dict(viewport={'width':w,'height':h},has_touch=False,is_mobile=False)
    elif dev.startswith('L'):
        base=dict(p.devices[dev[1:]]); vw=base['viewport']; base['viewport']={'width':vw['height'],'height':vw['width']}; base['screen']=dict(base['viewport']); cfg=base
    else: cfg=dict(p.devices[dev])
    cfg.pop('default_browser_type',None)
    b=p.chromium.launch(); ctx=b.new_context(accept_downloads=True,**cfg)
    pg=ctx.new_page();pg.set_default_timeout(5000);errs=[]
    pg.on('pageerror',lambda e:errs.append('PAGEERR '+str(e)));pg.on('console',lambda m:errs.append('CONSOLE '+m.text) if m.type=='error' else None)
    res=[]
    def ok(name,cond,info=''): res.append((name,bool(cond),str(info)[:160]))
    touch=cfg.get('has_touch',False)
    def tap(sel):
        el=pg.locator(sel).first;el.scroll_into_view_if_needed();(el.tap if touch else el.click)();pg.wait_for_timeout(60)
    def openall(): pg.evaluate("document.querySelectorAll('.fold').forEach(c=>c.classList.remove('collapsed'))")
    def kpi(l): return pg.evaluate("(l)=>{for(const k of document.querySelectorAll('.kpi')){if(k.innerText.includes(l))return k.querySelector('b').innerText}}",l)
    pg.goto(URL);pg.wait_for_timeout(500)
    W=pg.evaluate("innerWidth");mobile=W<=700
    ok('load/no errors',not errs,errs[:2])
    ok('rows=50',pg.evaluate("document.querySelectorAll('#tb tr').length")==50)
    ok('KPI total',kpi('Showing').startswith('287 /'),kpi('Showing'))
    ok('layout mode',(pg.evaluate("getComputedStyle(document.querySelector('#tb tr')).display")=='flex')==mobile,f'W={W}')
    ok('no h-overflow default',pg.evaluate("document.documentElement.scrollWidth")<=W)
    if mobile:
        ok('folds collapsed on mobile',pg.evaluate("document.querySelectorAll('.fold.collapsed').length")==7)
        ok('totop+msort visible',pg.evaluate("getComputedStyle(document.getElementById('totop')).display")=='block' and pg.evaluate("getComputedStyle(document.querySelector('.msort')).display")=='flex')
    openall();pg.wait_for_timeout(150)
    ok('no h-overflow all open',pg.evaluate("document.documentElement.scrollWidth")<=W,pg.evaluate("document.documentElement.scrollWidth"))
    if touch:
        small=pg.evaluate("[...document.querySelectorAll('button,select,input:not([type=checkbox]):not([type=range]):not([type=file])')].filter(e=>{const r=e.getBoundingClientRect();return r.width>0&&r.height>0&&r.height<38&&!e.closest('.cmp')&&!e.closest('#pgT,#pgB')}).length")
        ok('touch targets >=38px',small==0,small)
    pg.fill('#q','WIG');pg.wait_for_timeout(100);ok('search WIG',kpi('Showing').startswith(f"{exp['wig']} /"),(kpi('Showing'),exp['wig']));pg.fill('#q','')
    pg.fill('#rx','^[A-Z]{3} \\d{4}$');pg.wait_for_timeout(100);ok('regex filter',not kpi('Showing').startswith('287'))
    pg.fill('#rx','[');pg.wait_for_timeout(100);ok('invalid regex no crash',pg.evaluate("document.getElementById('rx').style.borderColor")!='' and not errs);pg.fill('#rx','')
    pg.fill('#q','ZZZZZZ');pg.wait_for_timeout(100);ok('empty state',pg.evaluate("document.querySelector('#tb').innerText").startswith('No plates'));pg.fill('#q','')
    pg.fill('#nm','boy');pg.wait_for_timeout(100);ok('name boy',kpi('Showing').startswith(f"{exp['boy']} /"),(kpi('Showing'),exp['boy']));ok('highlight mark',pg.evaluate("document.querySelectorAll('#tb mark').length")>0);pg.fill('#nm','')
    pg.select_option('#nmode','initc');pg.fill('#nm','David Williams');pg.wait_for_timeout(100);ok('initials DW none',kpi('Showing').startswith('0 /'));pg.fill('#nm','');pg.select_option('#nmode','contains')
    tap('#tpills .pill[data-t=Current]');ok('pill Current',kpi('Showing').startswith(f"{exp['cur']} /"),kpi('Showing'))
    tap('#tpills .pill[data-t=Suffix]');ok('pill Current+Suffix',kpi('Showing').startswith(f"{exp['cursuf']} /"))
    tap('#qpills .pill[data-q=""]');ok('clear pills',kpi('Showing').startswith('287 /'))
    tap('#qpills .pill[data-q=p1]');tap('#qpills .pill[data-q=u250]');ok('pills p1+u250',kpi('Showing').startswith(f"{exp['p1u250']} /"),(kpi('Showing'),exp['p1u250']))
    tap('#qpills .pill[data-q=p2]');ok('p1/p2 exclusive',pg.evaluate("document.querySelectorAll('#qpills .pill.on').length")==2);tap('#qpills .pill[data-q=""]')
    info=lambda:pg.inner_text('#pgB .info')
    ok('pager info',info().startswith('Showing 1–50 of 287'),info())
    tap('#pgB button:has-text("Next")');ok('next page',info().startswith('Showing 51–100'))
    pg.fill('#gopg','4');pg.press('#gopg','Enter');pg.wait_for_timeout(80);ok('go to page 4',info().startswith('Showing 151–200'))
    tap('#pgB button[title="Last page"]');ok('last page',info().startswith('Showing 251–287'))
    tap('#pgB button[title="First page"]');ok('first page',info().startswith('Showing 1–50'))
    pg.select_option('#ps','10');pg.wait_for_timeout(80);ok('10 per page',pg.evaluate("document.querySelectorAll('#tb tr').length")==10 and '29' in info())
    pg.select_option('#ps','All');pg.wait_for_timeout(80);ok('All rows',pg.evaluate("document.querySelectorAll('#tb tr').length")==287)
    pg.select_option('#ps','50');pg.wait_for_timeout(80)
    if not touch:
        pg.locator('body').click(position={'x':5,'y':5});pg.keyboard.press('ArrowRight');pg.wait_for_timeout(80);ok('arrow key paging',info().startswith('Showing 51'));pg.keyboard.press('ArrowLeft');pg.wait_for_timeout(80)
    first=lambda:pg.evaluate("document.querySelector('#tb tr .plate').innerText")
    if mobile: pg.select_option('#ms','price');pg.wait_for_timeout(80);a=first();tap('#md');b2=first()
    else: pg.select_option('#s1','price');pg.wait_for_timeout(80);a=first();pg.select_option('#d1','-1');pg.wait_for_timeout(80);b2=first()
    ok('sort price asc/desc',a==exp['cheapest'] and b2==exp['dearest'],(a,b2))
    pg.evaluate("document.getElementById('s1').value='alpha';document.getElementById('s1').dispatchEvent(new Event('input'))");pg.wait_for_timeout(80);ok('sort A-Z',first()=='1450 IL',first())
    pg.evaluate("document.getElementById('s1').value='n';document.getElementById('s1').dispatchEvent(new Event('input'))");pg.wait_for_timeout(80);ok('sort original',first()=='B443 AAE')
    pg.select_option('#rkm','short');pg.fill('#topn','5');tap('#rk-sel');ok('top5 select','5 selected' in pg.inner_text('#selinfo'),pg.inner_text('#selinfo'))
    tap('#b-clr');ok('clear selection','0 selected' in pg.inner_text('#selinfo'));pg.select_option('#rkm','price_lo')
    tap('#tb tr:nth-child(1) .rs');tap('#tb tr:nth-child(2) .rs');ok('select 2',pg.inner_text('#selinfo').startswith('2 selected'))
    tap('#b-add');ok('shortlist 2','2 plates' in pg.inner_text('#slinfo'),pg.inner_text('#slinfo'))
    ok('compare 2 cols',pg.evaluate("document.querySelectorAll('#cmp thead th').length")==3)
    ok('shortlist persisted',pg.evaluate("JSON.parse(localStorage.getItem('vrnfav')).length")==2)
    pg.select_option('#part','Part 1');pg.wait_for_timeout(100);ok('stats median £399',('£399' in pg.evaluate("document.getElementById('stk').innerText")));pg.select_option('#part','')
    bars=pg.locator('#hsvg .b');ok('hist bars',bars.count()>=10,bars.count())
    el=bars.nth(3);el.scroll_into_view_if_needed();(el.tap if touch else el.click)();pg.wait_for_timeout(120)
    ok('hist bar filters',pg.inner_text('#hinfo').endswith('£150–£199') and not kpi('Showing').startswith('287'),pg.inner_text('#hinfo')[-30:])
    tap('#hclr');ok('hist clear',kpi('Showing').startswith('287'))
    if not touch:
        b1=pg.locator('#hsvg .b').nth(2).bounding_box();b2_=pg.locator('#hsvg .b').nth(5).bounding_box()
        pg.mouse.move(b1['x']+b1['width']/2,b1['y']+b1['height']-2);pg.mouse.down();pg.mouse.move(b2_['x']+b2_['width']/2,b2_['y']+b2_['height']-2,steps=6);pg.mouse.up();pg.wait_for_timeout(120)
        ok('hist drag sets min/max',pg.input_value('#pmin')!='' and pg.input_value('#pmax')!='');tap('#hclr')
    t0=pg.evaluate("getComputedStyle(document.body).backgroundColor")
    tap('#theme');tap('#theme');pg.wait_for_timeout(400)
    ok('theme dark',pg.evaluate("document.documentElement.getAttribute('data-theme')")=='dark' and pg.evaluate("getComputedStyle(document.body).backgroundColor")!=t0)
    tap('#theme');pg.wait_for_timeout(100);ok('theme back to auto',pg.evaluate("document.documentElement.getAttribute('data-theme')")is None)
    pg.reload();pg.wait_for_timeout(700)
    ok('shortlist survives reload',pg.evaluate("document.querySelectorAll('.kpi')[8].innerText").startswith('2'))
    if heavy:
        openall()
        def dl(btn):
            with pg.expect_download() as di: (pg.locator(btn).tap if touch else pg.locator(btn).click)()
            return di.value
        pg.select_option('#escope','all')
        dc=dl('#x-csv');lines=[l for l in open(dc.path(),encoding='utf-8-sig').read().splitlines() if not l.startswith('#')]
        ok('csv rows',len(lines)==288,len(lines));ok('csv filename',dc.suggested_filename.endswith('-plates.csv') and '287' in dc.suggested_filename)
        dj=dl('#x-json');j=json.load(open(dj.path()));ok('json rows',j['count']==287 and len(j['rows'])==287)
        dx=dl('#x-xlsx');xp=os.path.join('/tmp','dl_'+re.sub(r'\W','_',dev)+'.xlsx');dx.save_as(xp);wb=openpyxl.load_workbook(xp);ws=wb['VRNs']
        ok('xlsx rows+sheets',ws.max_row==288 and wb.sheetnames==['VRNs','Filters','Statistics'],(ws.max_row,wb.sheetnames))
        tot=sum(v for v in (ws.cell(r,3).value for r in range(2,289)) if isinstance(v,(int,float)));ok('xlsx price sum',tot==sum(r['p'] for r in priced),tot)
        pg.evaluate("S.sel.clear();render()");tap('#tb tr:nth-child(1) .rs');tap('#tb tr:nth-child(2) .rs');pg.select_option('#escope','sel')
        dc=dl('#x-csv');lines=[l for l in open(dc.path(),encoding='utf-8-sig').read().splitlines() if not l.startswith('#')];ok('csv selected scope',len(lines)==3,len(lines))
        pg.select_option('#escope','all')
        pg.evaluate("window.print=()=>{window.__printed=true}");tap('#x-print');pg.wait_for_timeout(100)
        ok('print table built',pg.evaluate("document.querySelectorAll('#print tbody tr').length")==287 and pg.evaluate("!!window.__printed"))
        pg.fill('#nm','boy');pg.fill('#vname','Boys');tap('#v-save');pg.wait_for_timeout(100)
        ok('view saved',pg.evaluate("Object.keys(JSON.parse(localStorage.getItem('vrnviews'))).includes('Boys')"))
        pg.fill('#nm','');tap('#reset');ok('reset',kpi('Showing').startswith('287'))
        pg.select_option('#vlist','Boys');tap('#v-load');pg.wait_for_timeout(100);ok('view loaded',kpi('Showing').startswith(f"{exp['boy']} /"))
        link=pg.evaluate("(()=>{const st=getState();st.sel=[];return location.href.split('#')[0]+'#v='+b64(JSON.stringify(st))})()")
        pg2=ctx.new_page();pg2.goto(link);pg2.wait_for_timeout(400);ok('link restores view',pg2.evaluate("document.querySelector('.kpi b').innerText").startswith(f"{exp['boy']} /"));pg2.close()
        dv=dl('#v-exp');vj=json.load(open(dv.path()));ok('views export',('Boys' in vj['views']))
        pg.evaluate("document.getElementById('v-auto').checked=true;document.getElementById('v-auto').dispatchEvent(new Event('change'))");pg.wait_for_timeout(100)
        pg.reload();pg.wait_for_timeout(700);ok('reopen where left off',pg.evaluate("document.querySelector('.kpi b').innerText").startswith(f"{exp['boy']} /"))
    ctx2=b.new_context(**cfg);pgx=ctx2.new_page();e2=[];pgx.on('pageerror',lambda e:e2.append(str(e)))
    pgx.add_init_script("Object.defineProperty(window,'localStorage',{get(){throw new Error('blocked')}})")
    pgx.goto(URL);pgx.wait_for_timeout(400);ok('works with storage blocked',not e2 and pgx.evaluate("document.querySelectorAll('#tb tr').length")==50,e2[:1]);ctx2.close()
    ok('no errors overall',not errs,errs[:3])
    out={'device':dev,'width':W,'pass':sum(1 for r in res if r[1]),'total':len(res),'heavy':heavy,'fails':[r for r in res if not r[1]]}
    json.dump(out,open(os.path.join(OUT,'res_'+re.sub(r'\W','_',dev)+'.json'),'w'))
    b.close()
